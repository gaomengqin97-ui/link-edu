from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import FormulaRule

OUT = '临客LINK_9月8日交付计划.xlsx'
BLUE = '246BCE'
LIGHT_BLUE = 'EAF3FF'
PINK = 'E85AAD'
LIGHT_PINK = 'FCEAF6'
INK = '1F2937'
WHITE = 'FFFFFF'
GRAY = '6B7280'
GREEN = 'DCFCE7'
YELLOW = 'FEF3C7'

wb = Workbook()
ws = wb.active
ws.title = '项目总览'
daily = wb.create_sheet('每日排期')
accept = wb.create_sheet('分工与验收')

thin = Side(style='thin', color='D8E2F0')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def title(sheet, text, subtitle, end_col):
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    c = sheet.cell(1, 1, text)
    c.font = Font(name='Microsoft YaHei', size=20, bold=True, color=WHITE)
    c.fill = PatternFill('solid', fgColor=BLUE)
    c.alignment = Alignment(horizontal='left', vertical='center')
    sheet.row_dimensions[1].height = 34
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    c = sheet.cell(2, 1, subtitle)
    c.font = Font(name='Microsoft YaHei', size=10, color=GRAY, italic=True)
    c.alignment = Alignment(vertical='center')
    sheet.row_dimensions[2].height = 24

def header_row(sheet, row, values):
    for col, value in enumerate(values, 1):
        c = sheet.cell(row, col, value)
        c.font = Font(name='Microsoft YaHei', bold=True, color=WHITE)
        c.fill = PatternFill('solid', fgColor=PINK)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border
    sheet.row_dimensions[row].height = 30

def body(sheet, start_row, rows, fill_alt=True):
    for r, values in enumerate(rows, start_row):
        for col, value in enumerate(values, 1):
            c = sheet.cell(r, col, value)
            c.font = Font(name='Microsoft YaHei', size=10, color=INK)
            c.alignment = Alignment(vertical='top', wrap_text=True)
            c.border = border
            if fill_alt and (r - start_row) % 2:
                c.fill = PatternFill('solid', fgColor=LIGHT_BLUE)

title(ws, '临客 LINK｜9 月 8 日交付计划', '4 人团队 · 8 天冲刺 · 蓝粉版会议执行表', 6)
ws['A4'] = '交付日期'; ws['B4'] = '2026-09-08'
ws['C4'] = '团队人数'; ws['D4'] = 4
ws['E4'] = '策略'; ws['F4'] = 'P0 稳定闭环，P1 完成基础版，P2 轻量收口'
for c in ws[4]:
    c.font = Font(name='Microsoft YaHei', bold=c.column in (1,3,5), color=BLUE if c.column in (1,3,5) else INK)
    c.fill = PatternFill('solid', fgColor=LIGHT_PINK if c.column in (1,3,5) else WHITE)
    c.border = border
    c.alignment = Alignment(vertical='center', wrap_text=True)

header_row(ws, 6, ['优先级', '模块', '本期必须交付', '负责人', '验收标准', '不纳入本期'])
overview_rows = [
    ['P0', '教学训练核心闭环', '选阶段/完整训练；计时进度；摄像头显示；结束生成评课；失败提示', 'A（你）+ D', '登录→选课→训练→结束→评课连续跑通 3 次', '真实 ASR、TTS、数字人、CV 识别'],
    ['P1', 'AI 评课', '六维图表、详细报告、优点/改进、校正点重生成', 'C + D', '报告可查看、历史可查、校正后结果可保存', '外部大模型真实分析'],
    ['P1', '成长档案', '训练记录、30 天热力图、回放列表、训练日志', 'C + D', '页面显示数据库数据，日志可新增', '真实视频回放'],
    ['P1', '教学资源', '课程分阶段、资源分类、示范材料/来源说明', 'B（搞蒙圈）', '至少 6–8 门课程、资源可筛选和查看', '实时联网抓取资源'],
    ['P2', '个人中心', '资料编辑、等级、徽章、最近训练', 'C', '保存后刷新仍保留', '复杂设置中心'],
    ['P2', '帮助中心/视觉', '搜索 + FAQ + 客服占位；加载/空/错误态统一', 'A + D', '所有导航可进入，无白屏死链', '拖拽客服、全面动画重构'],
]
body(ws, 7, overview_rows)
tab = Table(displayName='PriorityTable', ref='A6:F12')
tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=False, showColumnStripes=False)
ws.add_table(tab)

title(daily, '每日排期｜8 月 31 日—9 月 8 日', '9 月 6 日完成完整彩排，9 月 7 日功能冻结', 8)
header_row(daily, 4, ['日期', '阶段目标', 'A｜你', 'B｜搞蒙圈', 'C｜业务模块', 'D｜技术集成', '当日交付物', '验收/闸门'])
daily_rows = [
    ['8/31', '冻结范围', '确定训练演示路径与视觉边界', '确定课程/资源目录', '确定评课/成长字段', '确认 API、数据库、演示数据', '范围清单、字段表、问题清单', '不再新增未评估功能'],
    ['9/1', 'P0 基础打通', '训练页主流程、全屏布局', '课程入口与 6–8 门课程', '评课数据结构', '训练创建/更新/完成接口', '训练可创建并保存', '完成一次闭环骨架'],
    ['9/2', 'P0 可演示', '摄像头、计时、阶段提示、异常态', '课程详情与训练入口', '接收训练完成结果', '登录、数据库、接口错误处理', 'P0 演示版本', '连续跑通登录→训练→评课'],
    ['9/3', 'P1 AI 评课', '评课页视觉统一', '补示范材料/来源', '六维图表、报告、校正重生成', '报告保存与历史查询', '完整评课页', '校正一次并成功保存'],
    ['9/4', 'P1 成长/资源', '空状态与响应式处理', '资源分类、筛选、详情', '热力图、回放、日志', '成长/资源接口联调', '成长档案、资源库基础版', '数据来自数据库'],
    ['9/5', 'P2 收口/联调', '统一按钮、加载、错误态', '校对课程/资源文案', '个人中心资料/徽章', '帮助中心搜索/客服占位', '所有导航可用', '无死链、无白屏'],
    ['9/6', '完整彩排', '按脚本跑 3 次', '修正资源展示问题', '修正报告/档案问题', '记录并清零阻塞问题', '演示脚本、问题清单', '三次连续演示成功'],
    ['9/7', '冻结交付', '只修阻塞视觉问题', '准备资源备份', '准备数据说明', '构建、部署、离线兜底、录屏', '发布包、备用视频、账号说明', '原则上不新增功能'],
    ['9/8', '正式交付', '现场演示支持', '现场资源说明', '现场业务说明', '启动与故障保障', '交付演示与说明', '完成最终验收'],
]
body(daily, 5, daily_rows)
tab = Table(displayName='DailyPlanTable', ref='A4:H13')
tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=False, showColumnStripes=False)
daily.add_table(tab)

title(accept, '分工与验收｜会议拍板版', '每个功能必须说明：真实能力、演示能力、验收方式', 7)
header_row(accept, 4, ['负责人', '工作流', '真实能力', '本期演示部分', '必须完成', '验收方式', '风险/依赖'])
accept_rows = [
    ['A｜你', '教学训练与视觉交互', '计时、进度、摄像头权限与画面显示、页面流程', '阶段提示、数字人/语音入口占位', '训练全屏、8/10 分钟模式、异常态', '连续完成 3 次训练演示', '浏览器摄像头权限；不承诺 CV/ASR'],
    ['B｜搞蒙圈', '课程与教学资源', '数据库课程/资源读取、分类展示', '示范内容为预置素材和来源说明', '6–8 门可选课程、资源分类筛选', '现场打开课程和资源详情', '素材版权、文案校对'],
    ['C｜业务模块', 'AI 评课与成长档案', '报告保存、历史查询、热力图/日志计算', '规则评分、固定模板摘要、校正加分', '六维报告、校正、30 天热力图、日志', '新建日志并校正一份报告', '无真实音视频分析'],
    ['D｜技术集成', '后端、联调、部署与 QA', 'Flask、SQLite、JWT、训练/评课/成长 API', '演示账号和种子数据', '接口稳定、启动说明、错误兜底、发布包', '健康检查 + 全链路彩排', '旧 Python 2.7 虚拟环境、启动脚本需规避'],
]
body(accept, 5, accept_rows)
tab = Table(displayName='OwnershipTable', ref='A4:G8')
tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=False, showColumnStripes=False)
accept.add_table(tab)

for sheet, widths in [(ws, [12, 20, 42, 16, 38, 30]), (daily, [11, 18, 31, 28, 30, 31, 28, 31]), (accept, [16, 22, 30, 31, 34, 30, 30])]:
    sheet.freeze_panes = 'A5' if sheet is not ws else 'A7'
    sheet.sheet_view.showGridLines = False
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None and cell.row not in (1, 2):
                cell.alignment = Alignment(vertical='top', wrap_text=True, horizontal=cell.alignment.horizontal)
    sheet.auto_filter.ref = sheet.tables[list(sheet.tables)[0]].ref

for sheet in (ws, daily, accept):
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.5
    sheet.page_margins.bottom = 0.5

wb.save(OUT)
print(OUT)
